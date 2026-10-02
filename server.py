"""Excel MCP Server: create / describe / edit / EDA / charts."""
import ast
import operator
import re
import statistics
import uuid
from collections import Counter
from pathlib import Path
from typing import Any, Optional

from mcp.server.fastmcp import FastMCP
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

mcp = FastMCP("excel-agent")

FILES_DIR = Path("files")
FILES_DIR.mkdir(exist_ok=True)


# ---------- helpers ----------
def _path(file_id: str) -> Path:
    p = FILES_DIR / f"{file_id}.xlsx"
    if not p.exists():
        raise ValueError(f"file_id '{file_id}' not found")
    return p


def _sheet(wb, sheet_name: Optional[str]):
    return wb[sheet_name] if sheet_name else wb.active


def _headers(ws) -> list:
    return [c.value for c in ws[1]]


def _style_header(ws):
    fill = PatternFill("solid", start_color="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 50)


# --- formula evaluation (so EDA can read formula cells) ---
_OPS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.Pow: operator.pow,
}


def _cell_value(ws, ref: str, depth: int = 0):
    """Value of a cell. If it holds a formula, calculate it."""
    v = ws[ref].value
    if isinstance(v, str) and v.startswith("="):
        return _eval_formula(v, ws, depth + 1)
    return v


def _eval_formula(formula: str, ws, depth: int = 0):
    """Evaluate simple arithmetic formulas like =C2/B2 or =(B2+C2)*2. Returns None if unsupported."""
    if depth > 10:
        return None
    expr = re.sub(
        r"\$?([A-Z]{1,3})\$?(\d+)",
        lambda m: f"_c('{m.group(1)}{m.group(2)}')",
        formula[1:].upper(),
    )

    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
            return _OPS[type(node.op)](ev(node.left), ev(node.right))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
            return -ev(node.operand)
        if isinstance(node, ast.Call) and getattr(node.func, "id", None) == "_c":
            v = _cell_value(ws, node.args[0].value, depth)
            if not isinstance(v, (int, float)):
                raise ValueError("not a number")
            return v
        raise ValueError("unsupported")

    try:
        return ev(ast.parse(expr, mode="eval"))
    except Exception:
        return None


# ---------- tools ----------
@mcp.tool()
def create_excel(
    columns: list[str],
    rows: list[list[Any]],
    sheet_name: str = "Sheet1",
) -> dict:
    """Create a new Excel file. Row 1 = headers. Returns file_id and path.
    Formulas can be written as strings like '=B2*C2'."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(columns)
    for r in rows:
        ws.append(r)
    _style_header(ws)
    _autofit(ws)
    ws.freeze_panes = "A2"

    file_id = uuid.uuid4().hex[:8]
    wb.save(FILES_DIR / f"{file_id}.xlsx")
    return {
        "file_id": file_id,
        "path": str(FILES_DIR / f"{file_id}.xlsx"),
        "rows": len(rows),
        "columns": columns,
    }


@mcp.tool()
def describe_excel(file_id: str, preview_rows: int = 5) -> dict:
    """Describe the file: sheets, headers, size, and first rows.
    Call this BEFORE editing so you know the structure."""
    wb = load_workbook(_path(file_id))
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = {
            "headers": _headers(ws),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "preview": [
                list(r)
                for r in ws.iter_rows(
                    min_row=2, max_row=min(ws.max_row, 1 + preview_rows), values_only=True
                )
            ],
        }
    return {"file_id": file_id, "sheets": out}


@mcp.tool()
def edit_excel(file_id: str, operations: list[dict], sheet_name: Optional[str] = None) -> dict:
    """Edit an existing Excel file. Header is row 1. `operations` is a list of:
      {"op":"set_cell","cell":"B2","value":10}
      {"op":"add_column","header":"Total","formula":"=B{row}*C{row}"}   # {row} is replaced per row
      {"op":"add_column","header":"Notes","values":["a","b"]}
      {"op":"delete_column","header":"Notes"}
      {"op":"rename_column","old":"A","new":"B"}
      {"op":"add_rows","rows":[[1,2,3],[4,5,6]]}
      {"op":"delete_rows","start":3,"count":2}
      {"op":"rename_sheet","new_name":"Data"}
    """
    path = _path(file_id)
    wb = load_workbook(path)
    ws = _sheet(wb, sheet_name)
    log = []

    for op in operations:
        kind = op.get("op")
        if kind == "set_cell":
            ws[op["cell"]] = op["value"]
        elif kind == "add_column":
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=op["header"])
            last = ws.max_row
            if "formula" in op:
                for r in range(2, last + 1):
                    ws.cell(row=r, column=col, value=op["formula"].replace("{row}", str(r)))
            elif "values" in op:
                for i, v in enumerate(op["values"]):
                    ws.cell(row=2 + i, column=col, value=v)
        elif kind == "delete_column":
            headers = _headers(ws)
            if op["header"] not in headers:
                raise ValueError(f"column '{op['header']}' not found. Available: {headers}")
            ws.delete_cols(headers.index(op["header"]) + 1)
        elif kind == "rename_column":
            headers = _headers(ws)
            if op["old"] not in headers:
                raise ValueError(f"column '{op['old']}' not found. Available: {headers}")
            ws.cell(row=1, column=headers.index(op["old"]) + 1, value=op["new"])
        elif kind == "add_rows":
            for r in op["rows"]:
                ws.append(r)
                new_row = ws.max_row
                prev = new_row - 1
                if prev < 2:
                    continue
                # extend formulas from the row above into columns the user didn't fill
                for col in range(len(r) + 1, ws.max_column + 1):
                    letter = get_column_letter(col)
                    above = ws.cell(row=prev, column=col).value
                    if isinstance(above, str) and above.startswith("="):
                        ws.cell(row=new_row, column=col).value = Translator(
                            above, origin=f"{letter}{prev}"
                        ).translate_formula(f"{letter}{new_row}")
        elif kind == "delete_rows":
            ws.delete_rows(op["start"], op.get("count", 1))
        elif kind == "rename_sheet":
            ws.title = op["new_name"]
        else:
            raise ValueError(f"unknown op: {kind}")
        log.append(kind)

    _style_header(ws)
    _autofit(ws)
    wb.save(path)
    return {"file_id": file_id, "applied": log, "headers": _headers(ws), "max_row": ws.max_row}


@mcp.tool()
def run_eda(file_id: str, sheet_name: Optional[str] = None) -> dict:
    """Run exploratory data analysis: shape, column types, missing values, duplicates,
    summary stats for numeric columns, correlations, and top categories."""
    wb = load_workbook(_path(file_id))
    ws = _sheet(wb, sheet_name)
    headers = _headers(ws)
    rows = [
        tuple(_cell_value(ws, c.coordinate) for c in row)
        for row in ws.iter_rows(min_row=2)
    ]
    cols = {h: [r[i] for r in rows] for i, h in enumerate(headers)}

    def is_num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool)

    numeric, categorical, missing = {}, {}, {}
    for h, vals in cols.items():
        present = [v for v in vals if v is not None]
        missing[h] = len(vals) - len(present)
        if present and all(is_num(v) for v in present):
            numeric[h] = present
        else:
            categorical[h] = present

    numeric_summary = {}
    for h, vals in numeric.items():
        numeric_summary[h] = {
            "count": len(vals),
            "mean": round(statistics.mean(vals), 2),
            "median": round(statistics.median(vals), 2),
            "std": round(statistics.stdev(vals), 2) if len(vals) > 1 else None,
            "min": round(min(vals), 2),
            "max": round(max(vals), 2),
        }

    correlations = {}
    names = list(numeric)
    for a in names:
        correlations[a] = {}
        for b in names:
            pairs = [(x, y) for x, y in zip(cols[a], cols[b]) if is_num(x) and is_num(y)]
            try:
                xs, ys = zip(*pairs)
                correlations[a][b] = round(statistics.correlation(xs, ys), 2)
            except Exception:
                correlations[a][b] = None

    return {
        "shape": [len(rows), len(headers)],
        "column_types": {h: ("numeric" if h in numeric else "text") for h in headers},
        "missing": missing,
        "duplicates": len(rows) - len(set(rows)),
        "numeric_summary": numeric_summary,
        "correlations": correlations if len(names) > 1 else {},
        "top_categories": {h: dict(Counter(v).most_common(5)) for h, v in categorical.items()},
    }


@mcp.tool()
def create_chart(
    file_id: str,
    chart_type: str,
    x_column: str,
    y_columns: list[str],
    title: str = "",
    sheet_name: Optional[str] = None,
) -> dict:
    """Add a native Excel chart to the sheet. chart_type: bar | line | pie.
    x_column = column used for category labels, y_columns = one or more numeric columns.
    (pie uses only the first y column)."""
    path = _path(file_id)
    wb = load_workbook(path)
    ws = _sheet(wb, sheet_name)
    headers = _headers(ws)

    for c in [x_column, *y_columns]:
        if c not in headers:
            raise ValueError(f"column '{c}' not found. Available: {headers}")

    chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
    if chart_type not in chart_classes:
        raise ValueError("chart_type must be one of: bar, line, pie")

    chart = chart_classes[chart_type]()
    chart.title = title
    last = ws.max_row

    ycols = y_columns[:1] if chart_type == "pie" else y_columns
    for y in ycols:
        idx = headers.index(y) + 1
        chart.add_data(Reference(ws, min_col=idx, min_row=1, max_row=last), titles_from_data=True)

    xi = headers.index(x_column) + 1
    chart.set_categories(Reference(ws, min_col=xi, min_row=2, max_row=last))
    chart.width, chart.height = 18, 9

    # place each new chart below the previous ones, to the right of the table
    anchor = f"{get_column_letter(ws.max_column + 2)}{2 + 20 * len(ws._charts)}"
    ws.add_chart(chart, anchor)
    wb.save(path)
    return {"file_id": file_id, "chart_type": chart_type, "title": title, "anchor": anchor}


if __name__ == "__main__":
    mcp.run()
